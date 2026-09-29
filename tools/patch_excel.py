# -*- coding: utf-8 -*-
"""ヒアリングフォームExcelにフォームと同じ変更を加える（2026-09-29の打ち合わせ内容）。

- 類型C: 01_申請内容 の C1-06〜C1-09（指定区域・指定空き家・居住歴・年齢）を削除
- 類型B: 05_補助金でやりたいこと の末尾に B5-21〜B5-28（報償費・広告宣伝費）を追加

変更後は tools/extract.py -> tools/build.py を実行すると、Excelの内容がそのままHTMLに反映される
（build.py の REMOVE_IDS / EXTRA_QUESTIONS は不要になるが、残しても二重には反映されない）。
実行前に必ず元ファイルのバックアップを取ること。
"""
import glob, os, sys
from copy import copy
import openpyxl
from openpyxl.formatting.formatting import ConditionalFormattingList
from openpyxl.worksheet.hyperlink import Hyperlink
from openpyxl.worksheet.datavalidation import DataValidation

sys.stdout.reconfigure(encoding="utf-8")
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SP = "'05_補助金でやりたいこと'"


def find(tag):
    return [x for x in glob.glob(os.path.join(ROOT, "*.xlsx")) if tag in os.path.basename(x)][0]


def copy_style(src, dst):
    dst._style = copy(src._style)


def remap_cf(ws, mapping):
    """条件付き書式の適用範囲を mapping(old_ref)->new_ref で付け替える"""
    old = ws.conditional_formatting
    new = ConditionalFormattingList()
    for cf in old:
        ref = str(cf.sqref)
        for rule in cf.rules:
            new.add(mapping.get(ref, ref), rule)
    ws.conditional_formatting = new


def patch_c():
    path = find("類型C")
    wb = openpyxl.load_workbook(path)
    ws = wb["01_申請内容"]
    if ws["A11"].value != "C1-06":
        print("C: 既に変更済みのためスキップ"); return
    # 1) 設問行 11〜14 を空行の書式にして消す
    blank = [ws.cell(15, c) for c in range(1, 7)]
    link_src = [ws.cell(16, c) for c in range(1, 7)]
    for r in range(11, 15):
        for c in range(1, 7):
            cell = ws.cell(r, c)
            cell.value = None
            copy_style(blank[c - 1], cell)
        ws.row_dimensions[r].height = None
    # 2) 「次へ進む」リンク行(16)を12行目へ
    for c in range(1, 7):
        copy_style(link_src[c - 1], ws.cell(12, c))
        ws.cell(12, c).value = link_src[c - 1].value
        ws.cell(16, c).value = None
        copy_style(blank[c - 1], ws.cell(16, c))
    ws.row_dimensions[12].height = 18.0
    ws.row_dimensions[16].height = None
    # 3) 入力規則: はい/いいえ系(E11:E14)を削除
    ws.data_validations.dataValidation = [dv for dv in ws.data_validations.dataValidation if str(dv.sqref) not in ("E11:E12", "E11 E12", "E13:E14", "E13 E14")]
    # 4) 条件付き書式・印刷範囲・注記
    remap_cf(ws, {"A6:F14": "A6:F10", "F6:F14": "F6:F10"})
    ws.print_area = "A1:E12"
    ws["A3"].value = ws["A3"].value.replace("指定区域〜年齢要件は補助金の上乗せに当てはまるかを確認する質問です。分からなければ「わからない」を選んでください。", "")
    # 5) はじめに の進み具合
    h = wb["はじめに"]
    for row in h.iter_rows():
        for c in row:
            if isinstance(c.value, str) and "'01_申請内容'!F6:F14" in c.value:
                c.value = c.value.replace("F6:F14", "F6:F10")
    # 6) 回答一覧: C1-06〜C1-09 の4行を削除して詰める
    g = wb["回答一覧（支援者用）"]
    rows = [n for n in range(3, g.max_row + 1) if str(g.cell(n, 1).value) in ("C1-06", "C1-07", "C1-08", "C1-09")]
    assert rows == list(range(rows[0], rows[0] + 4)), rows
    last = g.max_row
    g.move_range("A%d:F%d" % (rows[-1] + 1, last), rows=-4)
    wb.save(path)
    print("C: 保存", path)


def patch_b():
    path = find("類型B")
    wb = openpyxl.load_workbook(path)
    ws = wb["05_補助金でやりたいこと"]
    if ws["A26"].value == "B5-21":
        print("B: 既に変更済みのためスキップ"); return
    YOMI = "\n▼リストから選んでください"
    F_BLANK = '=IF(TRIM(SUBSTITUTE(INDEX($E:$E,ROW()),"　"," "))="","未入力","入力済")'
    def f_dep(dep_row):
        return ('=IF(TRIM(SUBSTITUTE(INDEX(%s!$E:$E,%d),"　"," "))="なし","対象外",'
                'IF(TRIM(SUBSTITUTE(INDEX($E:$E,ROW()),"　"," "))="","未入力","入力済"))') % (SP, dep_row)
    qs = [
        # id, 項目, おたずね, 記入例, 依存先行(None=なし), 選択式
        ("B5-21", "報償費（専門家等への謝礼）", "補助金を使って、専門家（コンサルタント・デザイナー・講師など）へ謝礼を支払う予定はありますか。ある場合は、依頼する内容と依頼先を教えてください。ない場合は「なし」（続く設問は記入不要になります）。", "例）商品開発のための料理研究家への監修依頼", None, False),
        ("B5-22", "報償費の目的・効果", "その依頼で、何を目的とし、今と比べて何がどのくらい良くなりますか。", "例）冷凍商品の味と歩留まりを改善し、新商品を2品開発する", 26, False),
        ("B5-23", "報償費の金額", "報償費のおおよその金額（税抜）を教えてください。", "例）30万円", 26, False),
        ("B5-24", "報償費の見積書", "報償費の見積書または依頼内容がわかる書類はありますか。（あり／依頼中／まだ）", "例）あり", 26, True),
        ("B5-25", "広告宣伝費", "補助金を使って、広告・宣伝（チラシ、パンフレット、看板、Web広告、ホームページ制作など）を行う予定はありますか。ある場合は、その内容を教えてください。ない場合は「なし」（続く設問は記入不要になります）。", "例）新商品を紹介するパンフレットの制作と、SNS広告", None, False),
        ("B5-26", "広告宣伝費の目的・効果", "その広告・宣伝で、誰に何を伝え、どのような効果を見込みますか。", "例）道内の旅行客に新商品を知ってもらい、通販の注文を月20件増やす", 30, False),
        ("B5-27", "広告宣伝費の金額", "広告宣伝費のおおよその金額（税抜）を教えてください。", "例）45万円", 30, False),
        ("B5-28", "広告宣伝費の見積書", "広告宣伝費の見積書はありますか。（あり／依頼中／まだ）", "例）依頼中", 30, True),
    ]
    # 1) 「次へ進む」リンク行(27)を35行目へ
    link_cells = [(ws.cell(27, c).value, copy(ws.cell(27, c)._style)) for c in range(1, 7)]
    hl = ws["E27"].hyperlink
    blank_style = copy(ws.cell(26, 1)._style)
    for c in range(1, 7):
        ws.cell(35, c).value = link_cells[c - 1][0]
        ws.cell(35, c)._style = link_cells[c - 1][1]
    ws["E27"].hyperlink = None
    ws["E35"].hyperlink = Hyperlink(ref="E35", location=hl.location, display=hl.display)
    ws.row_dimensions[35].height = 18.0
    ws.row_dimensions[34].height = None
    # 2) 設問行 26〜33 を追加（行25の書式をコピー）
    for i, (qid, label, prompt, example, dep, is_sel) in enumerate(qs):
        r = 26 + i
        for c in range(1, 7):
            ws.cell(r, c)._style = copy(ws.cell(25, c)._style)
        ws.cell(r, 5).value = None
        ws.cell(r, 1).value = qid
        ws.cell(r, 2).value = label
        ws.cell(r, 3).value = prompt + (YOMI if is_sel else "")
        ws.cell(r, 4).value = example
        ws.cell(r, 6).value = f_dep(dep) if dep else F_BLANK
        ws.row_dimensions[r].height = 66.0 if len(prompt) > 60 else 62.1
    for r in (27, 28, 29, 30, 31, 32, 33):
        pass
    for c in range(1, 7):  # 27行目は元リンク行の書式(太字等)が残っているので設問書式に揃え済み。34行目は空行書式
        ws.cell(34, c)._style = blank_style
    # 3) 入力規則(あり/依頼中/まだ) に E29, E33 を追加
    for dv in ws.data_validations.dataValidation:
        if dv.formula1 == '"あり,依頼中,まだ"':
            dv.add("E29"); dv.add("E33")
    # 4) 条件付き書式・印刷範囲・注記
    remap_cf(ws, {"A6:F25": "A6:F33", "F6:F25": "F6:F33"})
    ws.print_area = "A1:E35"
    ws["A3"].value = ws["A3"].value + " 専門家への謝礼（報償費）や広告宣伝費は、最後の設問（B5-21〜B5-28）でお伺いします。"
    # 5) はじめに の進み具合
    h = wb["はじめに"]
    for row in h.iter_rows():
        for c in row:
            if isinstance(c.value, str) and SP + "!F6:F25" in c.value:
                c.value = c.value.replace("F6:F25", "F6:F33")
    # 6) 回答一覧: B5-20 の直後に8行挿入
    g = wb["回答一覧（支援者用）"]
    at = [n for n in range(3, g.max_row + 1) if g.cell(n, 1).value == "B5-20"][0] + 1
    last = g.max_row
    g.move_range("A%d:F%d" % (at, last), rows=8)
    for i, (qid, label, prompt, example, dep, is_sel) in enumerate(qs):
        r = at + i
        src_row = row_src = at - 1
        for c in range(1, 7):
            g.cell(r, c)._style = copy(g.cell(row_src, c)._style)
        sheet_row = 26 + i
        g.cell(r, 1).value = qid
        g.cell(r, 2).value = "5. 補助金を使ってやりたいこと"
        g.cell(r, 3).value = label
        g.cell(r, 4).value = prompt
        g.cell(r, 5).value = '=IF(TRIM(INDEX(%s!$E:$E,%d))="","",INDEX(%s!$E:$E,%d))' % (SP, sheet_row, SP, sheet_row)
        g.cell(r, 6).value = "=INDEX(%s!$F:$F,%d)" % (SP, sheet_row)
        g.row_dimensions[r].height = g.row_dimensions[row_src].height
    wb.save(path)
    print("B: 保存", path)


if __name__ == "__main__":
    patch_c()
    patch_b()
